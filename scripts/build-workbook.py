#!/usr/bin/env python3
"""
Builds docs/SeedChain_Data_and_Test_Results.xlsx: the real datasets SeedChain holds, their sources and units,
the ingestion lineage, and a running log of every test run.

  # 1) record a test run (appends to docs/test-runs.json; nothing is ever edited or removed)
  python3 scripts/build-workbook.py log --name "Backend unit+integration" --env local --kind vitest --file out.json
  python3 scripts/build-workbook.py log --name "E2E smoke (production)"   --env production --kind e2e --file e2e.json
  python3 scripts/build-workbook.py log --name "Manual Chrome walkthrough" --env local --kind manual --file steps.json

  # 2) rebuild the workbook
  python3 scripts/build-workbook.py build --export-dir DIR --fao-zip Production_Crops_Livestock_E_All_Data_\\(Normalized\\).zip

DIR holds CSV exports of reference_statistics, ingestion_runs, raw_payloads and weather (psql \\copy ... csv header).
Requires: pip install openpyxl
"""
import argparse, csv, datetime as dt, glob, io, json, os, sys, zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOG = os.path.join(ROOT, "docs", "test-runs.json")
OUT = os.path.join(ROOT, "docs", "SeedChain_Data_and_Test_Results.xlsx")


def load_log():
    return json.load(open(LOG)) if os.path.exists(LOG) else []


def cmd_log(a):
    runs = load_log()
    data = json.load(open(a.file))
    items, at = [], dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    if a.kind == "vitest":
        for f in data["testResults"]:
            for t in f["assertionResults"]:
                items.append({"group": os.path.basename(f["name"]), "name": t["fullName"], "status": "PASS" if t["status"] == "passed" else "FAIL", "detail": f'{round(t.get("duration") or 0)} ms'})
        at = dt.datetime.fromtimestamp(data["startTime"] / 1000, dt.timezone.utc).isoformat(timespec="seconds")
    elif a.kind == "e2e":
        items = data["results"]; at = data["at"]
    elif a.kind == "manual":
        items = data
    run = {"id": f"R{len(runs) + 1:03d}", "at": at, "name": a.name, "environment": a.env, "kind": a.kind, "notes": a.notes or "", "items": items}
    runs.append(run)
    json.dump(runs, open(LOG, "w"), indent=1)
    p = sum(1 for i in items if i["status"] == "PASS")
    print(f'{run["id"]}: {a.name} [{a.env}] {p}/{len(items)} passed')


def rd(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def cmd_build(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    HEAD = PatternFill("solid", fgColor="1F3D2B")
    PASS, FAIL = PatternFill("solid", fgColor="D7F0DC"), PatternFill("solid", fgColor="F8D7D7")

    def sheet(title, header, rows, widths=None, status_col=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(vertical="center", wrap_text=True)
        n = 0
        for r in rows:
            ws.append(r); n += 1
            if status_col is not None:
                c = ws.cell(row=ws.max_row, column=status_col + 1)
                c.fill = PASS if c.value == "PASS" else FAIL if c.value == "FAIL" else PatternFill()
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            w = (widths or {}).get(i - 1) or min(max(len(str(h)) + 2, 12), 40)
            ws.column_dimensions[get_column_letter(i)].width = w
        return n

    today = dt.date.today().isoformat()
    # --- FAOSTAT world potatoes (streamed from the official bulk file) ---
    fao = []
    with zipfile.ZipFile(a.fao_zip) as z:
        name = next(n for n in z.namelist() if n.endswith("All_Data_(Normalized).csv"))
        with z.open(name) as raw:
            for x in csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8", errors="replace")):
                if x["Item Code"] == "116" and x["Element Code"] in ("5312", "5412", "5510"):
                    ac = int(x["Area Code"])
                    fao.append([x["Area"], ac, "Region/aggregate" if ac >= 5000 else "Country", x["Item"], x["Element"], int(x["Year"]), x["Unit"], float(x["Value"] or 0), x["Flag"]])
    fao.sort(key=lambda r: (r[0], r[4], r[5]))
    countries = len({r[0] for r in fao if r[2] == "Country"})

    runs = load_log()
    total_tests = sum(len(r["items"]) for r in runs)
    readme = [
        ["Workbook", "SeedChain: real datasets and test results"], ["Built on", today],
        ["Purpose", "Everything SeedChain shows as external data, with its source, units and lineage, plus the log of every test run."],
        ["Rule", "No value here is invented. Each row comes from the named publisher; rows that could not be fetched are absent, not estimated."],
        ["", ""],
        ["Sheet", "What it holds"],
        ["Source_Registry", "Every external source we know of: publisher, URL, geography, update frequency, units, licence, and whether it is integrated."],
        ["PAU_Punjab_Districts", "Punjab Dept. of Horticulture potato table, republished by PAU: district area, production, yield for the latest published year."],
        ["FAOSTAT_India", "India potato area, production, yield 1961-2024 (FAOSTAT QCL), as stored in SeedChain's database."],
        [f"FAOSTAT_World_Potato", f"{len(fao):,} FAOSTAT records: potatoes, all {countries} countries plus regions, area/yield/production, 1961-2024. Official bulk file, filtered to item 116."],
        ["Units_Resources", "Unit definitions and the official links to every resource used."],
        ["Ingestion_Runs", "Every fetch the pipeline made: status, counts, endpoint, timings. Failures are shown, not hidden."],
        ["Raw_Payloads", "Fingerprint (SHA-256) and size of each untouched raw download kept for audit."],
        ["Mandi_Prices", "Daily potato mandi prices (INR per quintal): min, max, modal, arrivals. Shows NO DATA YET until the data.gov.in sync or a CSV import succeeds."],
        ["Weather_Observations", "Real Open-Meteo readings stored for farms that have coordinates (positions rounded to 0.01 degrees)."],
        ["Test_Runs", "One row per test run (this log only grows)."],
        ["Test_Results", f"One row per individual check or test across all runs ({total_tests:,} so far)."],
        ["", ""],
        ["Not in this workbook", "Mandi prices (see the Mandi_Prices sheet; empty until a key or CSV import), AGMARKNET, DES, NHB, ICAR-CPRI, IMD: not integrated, so no rows exist."],
        ["Note on FAOSTAT scope", "FAOSTAT is national/regional only. It has no state or district data; Punjab district figures come only from the PAU sheet."],
        ["Test data", "Accounts and lots named [E2E TEST] / [PILOT TEST] are test records created by the test scripts, not real farmers."],
    ]
    sheet("README", ["Item", "Detail"], readme, {0: 24, 1: 120}, first=True)
    for row in wb["README"].iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    reg = []
    for p in sorted(glob.glob(os.path.join(ROOT, "data", "source_registry", "*.json"))):
        d = json.load(open(p))
        reg.append([d["id"], d["name"], d["organization"], d["url"], d["dataset"], d["geography"], d["frequency"], d.get("units", ""), d.get("licence", ""), d["dataClass"], d["integration"], d.get("integrationNote", "")])
    sheet("Source_Registry", ["ID", "Name", "Publisher", "URL", "Dataset", "Geography", "Update frequency", "Units", "Licence / terms", "Data class", "Integration", "Verification note"], reg, {0: 22, 1: 40, 2: 36, 3: 44, 4: 60, 5: 28, 6: 30, 7: 36, 8: 40, 9: 16, 10: 16, 11: 60})

    stats = rd(os.path.join(a.export_dir, "reference_statistics.csv"))
    pau = [[s["geography_level"], s["geography_name"], s["metric"], s["unit"], s["period_label"], float(s["value"]), s["citation"], s["retrieved_at"][:19]] for s in stats if s["source_id"] == "pau_potato_punjab"]
    sheet("PAU_Punjab_Districts", ["Level", "Geography", "Metric", "Unit", "Period", "Value", "Citation", "Retrieved (UTC)"], pau, {1: 22, 6: 90, 7: 20})
    fi = [[s["geography_name"], s["metric"], s["unit"], int(s["year_start"]), float(s["value"]), s["retrieved_at"][:19]] for s in stats if s["source_id"] == "faostat_potato_india"]
    sheet("FAOSTAT_India", ["Geography", "Metric", "Unit", "Year", "Value", "Retrieved (UTC)"], fi, {0: 14, 5: 20})
    sheet("FAOSTAT_World_Potato", ["Area", "Area code", "Level", "Item", "Element", "Year", "Unit", "Value", "FAO flag"], fao, {0: 30, 2: 18, 4: 16})

    units = [["ha", "hectare", "Area harvested / cultivated area (10,000 m2)"], ["t", "tonne", "Production, metric tonnes (1,000 kg)"], ["kg/ha", "kilograms per hectare", "FAOSTAT yield unit"], ["q/ha", "quintals per hectare", "PAU yield unit (1 quintal = 100 kg); 1 q/ha = 100 kg/ha"], ["INR/quintal", "rupees per quintal", "Mandi price unit used by data.gov.in"], ["FAO flag", "A=official, E=estimated, I=imputed, M=missing, X=international figure", "FAO flags are kept as published"]]
    for d in sorted(glob.glob(os.path.join(ROOT, "data", "source_registry", "*.json"))):
        j = json.load(open(d)); units.append(["LINK", j["name"], j["url"]])
    sheet("Units_Resources", ["Unit / type", "Meaning / resource", "Detail / URL"], units, {0: 16, 1: 70, 2: 100})

    runs_csv = rd(os.path.join(a.export_dir, "ingestion_runs.csv"))
    sheet("Ingestion_Runs", ["Run ID", "Source", "Status", "Fetched", "Accepted", "Rejected", "Parser version", "Endpoint", "Error", "Started (UTC)", "Finished (UTC)"], [[r["id"], r["source"], r["status"], int(r["fetched_count"] or 0), int(r["record_count"] or 0), int(r["rejected_count"] or 0), r["processing_version"], r["source_endpoint"].split("?")[0], r["error"], r["started_at"][:19], r["finished_at"][:19]] for r in runs_csv], {0: 38, 1: 24, 7: 70, 8: 40, 9: 20, 10: 20}, status_col=None)
    sheet("Raw_Payloads", ["Run ID", "Source", "Content type", "Bytes", "SHA-256", "Retrieved (UTC)"], [[r["run_id"], r["source_id"], r["content_type"], int(r["byte_size"]), r["sha256"], r["retrieved_at"][:19]] for r in rd(os.path.join(a.export_dir, "raw_payloads.csv"))], {0: 38, 1: 24, 2: 24, 4: 70, 5: 20})
    wx = rd(os.path.join(a.export_dir, "weather.csv"))
    sheet("Weather_Observations", ["Source", "Observed (UTC)", "Retrieved (UTC)", "Lat (0.01)", "Lon (0.01)", "Temp C", "Rain mm", "Humidity %", "WMO code", "Condition"], [[w["source"], w["observation_time"][:19], w["retrieved_at"][:19], float(w["latitude"]), float(w["longitude"]), *(float(w[k]) if w[k] else None for k in ("temperature_c", "precipitation_mm", "humidity_pct")), w["weather_code"], w["condition"]] for w in wx], {1: 20, 2: 20})

    mp = os.path.join(a.export_dir, "mandi.csv")
    mandi = rd(mp) if os.path.exists(mp) else []
    mrows = [[m["observation_date"], m["state"], m["district"], m["market"], m["commodity"], m["variety"], m["grade"], float(m["arrival_quantity_tonnes"]) if m["arrival_quantity_tonnes"] else None, float(m["min_price"]) if m["min_price"] else None, float(m["max_price"]) if m["max_price"] else None, float(m["modal_price"]) if m["modal_price"] else None, m["price_unit"], m["source"], m["retrieved_at"][:19]] for m in mandi]
    if not mrows:
        mrows = [["NO DATA YET", "", "", "", "", "", "", None, None, None, None, "", "No successful mandi sync. Needs DATA_GOV_IN_API_KEY or a CSV import (see docs/DATA_SOURCES.md). Nothing is estimated.", ""]]
    sheet("Mandi_Prices", ["Observation date", "State", "District", "Market", "Commodity", "Variety", "Grade", "Arrivals (t)", "Min price", "Max price", "Modal price", "Unit", "Source / status", "Retrieved (UTC)"], mrows, {0: 16, 3: 28, 12: 60, 13: 20})
    tr, td = [], []
    for r in runs:
        p = sum(1 for i in r["items"] if i["status"] == "PASS"); f = len(r["items"]) - p
        tr.append([r["id"], r["at"].replace("T", " ")[:19], r["name"], r["environment"], r["kind"], len(r["items"]), p, f, "PASS" if f == 0 else "FAIL", r["notes"]])
        for i in r["items"]:
            td.append([r["id"], r["environment"], i.get("group", ""), i["name"], i["status"], i.get("detail", "")])
    sheet("Test_Runs", ["Run", "When (UTC)", "Suite", "Environment", "Type", "Checks", "Passed", "Failed", "Result", "Notes"], tr, {1: 20, 2: 44, 9: 90}, status_col=8)
    sheet("Test_Results", ["Run", "Environment", "Group / file", "Check", "Result", "Detail"], td, {2: 30, 3: 90, 5: 60}, status_col=4)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}: FAOSTAT world {len(fao):,} rows, PAU {len(pau)}, India {len(fi)}, runs {len(runs_csv)}, test results {len(td):,}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    l = sub.add_parser("log"); l.add_argument("--name", required=True); l.add_argument("--env", required=True); l.add_argument("--kind", required=True, choices=["vitest", "e2e", "manual"]); l.add_argument("--file", required=True); l.add_argument("--notes")
    b = sub.add_parser("build"); b.add_argument("--export-dir", required=True); b.add_argument("--fao-zip", required=True)
    a = ap.parse_args()
    {"log": cmd_log, "build": cmd_build}[a.cmd](a)
